"""Build the editable Study RAG concept and roadmap presentation."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = "study-rag-vision.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

NAVY = "101923"
NAVY2 = "172431"
PANEL = "1D2D3A"
PANEL2 = "243744"
CREAM = "F5F2EA"
PAPER = "EAE6DC"
INK = "18232B"
MUTED = "91A2AA"
MUTED_DARK = "63727A"
TEAL = "5CE0C1"
CORAL = "FF795F"
GOLD = "F3C76A"
BLUE = "82B6FF"
WHITE = "FFFFFF"
FONT = "Aptos"
FONT_HEAD = "Aptos Display"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def rect(slide, x, y, w, h, fill, radius=True, line=None, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def ellipse(slide, x, y, w, h, fill, line=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.2, dash=None):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    return shape


def text(slide, x, y, w, h, value, size=14, color=INK, bold=False,
         font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0, italic=False, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, part in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.line_spacing = spacing
        run = p.add_run()
        run.text = part
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
    return box


def label(slide, x, y, value, color=TEAL, w=3.2):
    text(slide, x, y, w, .22, value.upper(), 9, color, True, FONT, spacing=1)


def header(slide, section, title_value, subtitle=None, dark=True):
    fg = CREAM if dark else INK
    muted = MUTED if dark else MUTED_DARK
    label(slide, .72, .48, section)
    text(slide, .72, .88, 11.8, .64, title_value, 28, fg, True, FONT_HEAD)
    if subtitle:
        text(slide, .74, 1.58, 11.7, .38, subtitle, 12, muted)


def footer(slide, n, dark=True):
    color = "51616A" if dark else "B6B2A9"
    line(slide, .72, 7.08, 12.61, 7.08, color, .65)
    text(slide, .73, 7.14, 5, .18, "STUDY RAG  /  CONCEPT + ROADMAP", 8, MUTED if dark else MUTED_DARK, True)
    text(slide, 12.0, 7.12, .55, .2, f"{n:02d}", 9, TEAL if dark else INK, True, align=PP_ALIGN.RIGHT)


def new_slide(bg=NAVY):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=1.5):
    ln = line(slide, x1, y1, x2, y2, color, width)
    ln.line.end_arrowhead = True
    return ln


# 01 — Cover
s = new_slide()
ellipse(s, 8.7, .35, 4.1, 4.1, PANEL, transparency=8)
ellipse(s, 9.3, .95, 2.9, 2.9, NAVY, line=PANEL2)
ellipse(s, 10.0, 1.65, 1.5, 1.5, TEAL, transparency=5)
ellipse(s, 10.43, 2.08, .64, .64, NAVY)
label(s, .85, .74, "A local-first study assistant", TEAL, 4.5)
text(s, .82, 1.45, 7.3, 1.45, "STUDY\nRAG", 49, CREAM, True, FONT_HEAD, spacing=.83)
text(s, .9, 4.08, 6.6, .82, "From scattered course files to\ntraceable answers — including diagrams.", 20, "D7E0DE", False, FONT_HEAD, spacing=1.05)
rect(s, .91, 5.45, 2.35, .38, PANEL2)
text(s, 1.08, 5.54, 2.0, .16, "PROJECT VISION  /  2026", 8, TEAL, True)
# Floating document cards
rect(s, 8.4, 4.25, 3.5, 1.45, PANEL, line=PANEL2)
text(s, 8.72, 4.52, 2.65, .3, "Evidence, not just answers", 14, CREAM, True)
text(s, 8.72, 4.95, 2.85, .42, "Source · page · passage · figure", 10, MUTED)
ellipse(s, 11.25, 4.62, .34, .34, TEAL)
text(s, 11.34, 4.69, .2, .12, "✓", 9, NAVY, True)
text(s, 8.72, 6.36, 3.6, .22, "A ROADMAP FOR VISUAL DOCUMENT RAG", 9, MUTED, True)
footer(s, 1)


# 02 — Why visuals matter
s = new_slide(CREAM)
header(s, "01  /  THE OPPORTUNITY", "A page is more than its text layer", "Study material carries meaning in labels, layout, tables, and the links between visual elements.", False)
# Mock lecture page
rect(s, .82, 2.35, 4.35, 4.1, WHITE, line="D8D3C8")
text(s, 1.15, 2.65, 3.5, .26, "CELL SIGNALING", 10, MUTED_DARK, True)
text(s, 1.15, 3.03, 3.45, .42, "A pathway with a story", 19, INK, True, FONT_HEAD)
line(s, 1.15, 3.65, 4.83, 3.65, "DED9CF", .8)
rect(s, 1.18, 4.02, 1.14, .62, "DCEBE7")
text(s, 1.27, 4.2, .98, .16, "SIGNAL", 9, INK, True, align=PP_ALIGN.CENTER)
rect(s, 3.75, 4.02, 1.14, .62, "FCE7DF")
text(s, 3.84, 4.2, .98, .16, "RESPONSE", 9, INK, True, align=PP_ALIGN.CENTER)
arrow(s, 2.39, 4.33, 3.65, 4.33, CORAL, 2)
ellipse(s, 2.83, 4.19, .3, .3, GOLD)
text(s, 1.18, 4.94, 3.55, .9, "The arrow direction and intermediate step carry the idea — they are not ordinary prose.", 12, INK, False, spacing=1.05)
text(s, 1.18, 6.07, 3.3, .16, "ILLUSTRATIVE PAGE — NOT SOURCE MATERIAL", 7, MUTED_DARK, True)
# Insight panel
rect(s, 5.65, 2.35, 6.83, 1.18, "E4E0D6")
text(s, 5.98, 2.63, 1.6, .38, "TEXT", 14, INK, True, FONT_HEAD)
text(s, 7.2, 2.62, 4.75, .46, "What the page says", 15, MUTED_DARK, False, FONT_HEAD)
rect(s, 5.65, 3.78, 6.83, 1.18, "DDEAE5")
text(s, 5.98, 4.06, 1.6, .38, "LAYOUT", 14, INK, True, FONT_HEAD)
text(s, 7.2, 4.04, 4.8, .46, "How the ideas connect", 15, "397C6B", False, FONT_HEAD)
rect(s, 5.65, 5.2, 6.83, 1.18, "F7E4D9")
text(s, 5.98, 5.48, 1.6, .38, "EVIDENCE", 14, INK, True, FONT_HEAD)
text(s, 7.2, 5.45, 4.9, .48, "Where to verify the answer", 15, "A35442", False, FONT_HEAD)
footer(s, 2, False)


# 03 — Current pipeline
s = new_slide()
header(s, "02  /  CURRENT BASELINE", "A useful foundation — with a text-shaped lens", "The current prototype is intentionally small: local ingestion, persistent vectors, and cited search chunks.")
steps = [
    ("01", "COURSE FILES", "PDF · TXT · MD", BLUE),
    ("02", "EXTRACT + CHUNK", "pypdf · ~800 chars", GOLD),
    ("03", "EMBED", "MiniLM · local", TEAL),
    ("04", "RETRIEVE", "Chroma · top-k", CORAL),
    ("05", "RETURN", "MCP · citations", BLUE),
]
xs = [.82, 3.25, 5.68, 8.11, 10.54]
for i, (num, title, desc, accent) in enumerate(steps):
    rect(s, xs[i], 2.7, 1.95, 1.64, PANEL, line=PANEL2)
    ellipse(s, xs[i]+.17, 2.91, .34, .34, accent)
    text(s, xs[i]+.22, 2.99, .22, .12, num[-1], 8, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, xs[i]+.18, 3.46, 1.62, .22, title, 9, CREAM, True)
    text(s, xs[i]+.18, 3.79, 1.6, .3, desc, 9, MUTED)
    if i < len(steps)-1:
        arrow(s, xs[i]+2.02, 3.51, xs[i+1]-.13, 3.51, "62747D", 1.2)
# lower insight band
rect(s, .82, 4.92, 11.67, 1.23, "1A2935")
label(s, 1.12, 5.18, "What already works", TEAL, 2.3)
text(s, 1.12, 5.55, 3.2, .3, "Simple · local · inspectable", 15, CREAM, True, FONT_HEAD)
line(s, 5.1, 5.16, 5.1, 5.9, "40535D", .9)
label(s, 5.48, 5.18, "Current blind spot", CORAL, 2.2)
text(s, 5.48, 5.55, 6.4, .3, "The index sees extracted words, not diagram meaning.", 15, CREAM, True, FONT_HEAD)
footer(s, 3)


# 04 — Failure mode
s = new_slide(CREAM)
header(s, "03  /  THE GAP", "A diagram query needs more than OCR", "Labels are clues. The visual relationship between labels can be the answer.", False)
# Left side: diagram example
rect(s, .83, 2.34, 5.54, 3.98, WHITE, line="D8D3C8")
label(s, 1.13, 2.62, "A FLOWCHART ON THE PAGE", "A35442", 3)
rect(s, 1.24, 3.43, 1.28, .72, "DCEBE7")
text(s, 1.32, 3.68, 1.1, .18, "INPUT", 10, INK, True, align=PP_ALIGN.CENTER)
rect(s, 3.05, 3.43, 1.28, .72, "F7E4D9")
text(s, 3.13, 3.68, 1.1, .18, "FILTER", 10, INK, True, align=PP_ALIGN.CENTER)
rect(s, 4.85, 3.43, 1.12, .72, "E8E3F4")
text(s, 4.93, 3.68, .95, .18, "OUTPUT", 10, INK, True, align=PP_ALIGN.CENTER)
arrow(s, 2.56, 3.79, 2.94, 3.79, CORAL, 2)
arrow(s, 4.37, 3.79, 4.73, 3.79, CORAL, 2)
text(s, 1.22, 4.58, 4.6, .7, "Question: What must happen before output?\nOCR returns three labels. It may miss the sequence.", 12, INK, False, spacing=1.12)
text(s, 1.22, 5.77, 4.7, .2, "ILLUSTRATIVE DIAGRAM — NOT SOURCE MATERIAL", 7, MUTED_DARK, True)
# Right side, two result styles
rect(s, 6.75, 2.34, 5.72, 1.68, NAVY)
label(s, 7.08, 2.64, "TEXT-ONLY HIT", GOLD, 2)
text(s, 7.08, 3.03, 4.96, .52, "“Input · Filter · Output”\nPage unknown. Relationship unknown.", 14, CREAM, False, FONT_HEAD, spacing=1.06)
rect(s, 6.75, 4.28, 5.72, 2.04, "DDEAE5")
label(s, 7.08, 4.58, "EVIDENCE-AWARE HIT", "397C6B", 2.8)
text(s, 7.08, 4.98, 4.92, .66, "“Input flows through a filter before producing output.”", 15, INK, True, FONT_HEAD, spacing=1.0)
text(s, 7.08, 5.82, 4.9, .2, "Source PDF  ·  Page 12  ·  Figure 3", 9, "397C6B", True)
footer(s, 4, False)


# 05 — Target architecture
s = new_slide()
header(s, "04  /  TARGET EXPERIENCE", "One source. Three evidence paths. One citation.", "Structure the document once; retrieve the evidence type that fits the question.")
# input source card
rect(s, .82, 2.65, 2.05, 2.1, PANEL, line=PANEL2)
text(s, 1.08, 2.98, 1.45, .24, "SOURCE PDF", 10, BLUE, True)
rect(s, 1.15, 3.46, 1.18, .96, "293B47")
text(s, 1.33, 3.72, .82, .36, "PAGE\n12", 11, CREAM, True, align=PP_ALIGN.CENTER, spacing=.9)
text(s, 1.08, 4.48, 1.55, .18, "immutable input", 9, MUTED)
arrow(s, 2.95, 3.7, 3.54, 3.7, TEAL, 1.6)
# parser card
rect(s, 3.68, 2.65, 2.12, 2.1, "20313C", line="314652")
label(s, 3.97, 2.96, "PARSE", TEAL, 1.4)
text(s, 3.97, 3.38, 1.6, .75, "Layout\nOCR · figures\nPage map", 13, CREAM, True, FONT_HEAD, spacing=1.08)
arrow(s, 5.9, 3.7, 6.45, 3.7, TEAL, 1.6)
# three branches
line(s, 6.62, 3.7, 6.62, 2.68, "51636C", 1)
line(s, 6.62, 2.68, 7.08, 2.68, "51636C", 1)
line(s, 6.62, 3.7, 7.08, 3.7, "51636C", 1)
line(s, 6.62, 3.7, 6.62, 4.72, "51636C", 1)
line(s, 6.62, 4.72, 7.08, 4.72, "51636C", 1)
for y, accent, title_v, desc in [
    (2.3, BLUE, "TEXT", "Headings · paragraphs · tables"),
    (3.32, GOLD, "FIGURE", "Image · caption · generated description"),
    (4.34, CORAL, "PAGE IMAGE", "Original visual evidence"),
]:
    rect(s, 7.1, y, 3.05, .78, PANEL, line=PANEL2)
    ellipse(s, 7.29, y+.22, .22, .22, accent)
    text(s, 7.68, y+.12, 1.95, .2, title_v, 9, accent, True)
    text(s, 7.68, y+.39, 2.28, .22, desc, 8, CREAM)
    arrow(s, 10.22, y+.38, 10.56, y+.38, "60727B", 1.1)
# retrieval / evidence bundle
rect(s, 10.68, 3.02, 1.82, 1.4, "DDEAE5")
text(s, 10.93, 3.3, 1.3, .22, "EVIDENCE", 10, "397C6B", True)
text(s, 10.93, 3.68, 1.32, .5, "Answer\n+ page link", 13, INK, True, FONT_HEAD, spacing=1)
rect(s, 3.68, 5.73, 8.82, .65, "1A2935")
text(s, 3.97, 5.95, 8.1, .2, "Every record carries source · page · content type · provenance", 11, CREAM, True)
footer(s, 5)


# 06 — Diagram strategy
s = new_slide(CREAM)
header(s, "05  /  DIAGRAMS", "Layer the signals. Keep the image as ground truth.", "A caption makes a figure searchable; the original page lets a student verify it.", False)
cols = [(.82, BLUE, "01  READ LABELS", "OCR / extracted text", "Find printed terms and labels.", "Useful for named parts; weak on arrows and spatial meaning."),
        (4.98, "397C6B", "02  DESCRIBE", "Generated caption", "Translate visible structure into searchable prose.", "Useful for relationships; mark it generated and retain uncertainty."),
        (9.14, "A35442", "03  INSPECT", "Page / figure image", "Show the source visual alongside its citation.", "Useful for verification and questions that need visual reasoning.")]
for x, accent, title_v, subtitle, desc, note in cols:
    rect(s, x, 2.54, 3.37, 3.79, WHITE, line="D8D3C8")
    ellipse(s, x+.28, 2.83, .35, .35, accent)
    text(s, x+.8, 2.87, 2.23, .2, title_v, 10, accent, True)
    text(s, x+.28, 3.42, 2.8, .24, subtitle, 15, INK, True, FONT_HEAD)
    # distinct icon-like native vector
    if "LABELS" in title_v:
        rect(s, x+.36, 3.98, 1.08, .58, "E9F0F9")
        text(s, x+.48, 4.18, .82, .14, "A → B", 12, INK, True, align=PP_ALIGN.CENTER)
        line(s, x+1.7, 4.04, x+2.55, 4.04, "C9D4DE", 1)
        line(s, x+1.7, 4.27, x+2.8, 4.27, "C9D4DE", 1)
    elif "DESCRIBE" in title_v:
        rect(s, x+.36, 3.98, 2.72, .58, "E6F0EC")
        text(s, x+.53, 4.14, 2.42, .25, "A activates B, then C", 10, INK, True, align=PP_ALIGN.CENTER)
    else:
        rect(s, x+.36, 3.98, 1.08, .58, "F7E8E3")
        text(s, x+.49, 4.18, .82, .14, "PAGE 12", 10, INK, True, align=PP_ALIGN.CENTER)
        rect(s, x+1.68, 3.98, 1.4, .58, "F1E5E1")
        ellipse(s, x+2.18, 4.11, .3, .3, accent)
    text(s, x+.28, 4.87, 2.78, .53, desc, 11, INK, True, spacing=1.05)
    text(s, x+.28, 5.59, 2.82, .47, note, 9, MUTED_DARK, False, spacing=1.05)
rect(s, .82, 6.54, 11.69, .3, "E7E2D8")
text(s, 1.04, 6.61, 11.1, .14, "PROVENANCE RULE   Source text ≠ OCR labels ≠ model-generated description", 8, INK, True)
footer(s, 6, False)


# 07 — Roadmap
s = new_slide()
header(s, "06  /  BUILD SEQUENCE", "Ship the simplest useful visual RAG first", "Prove each layer against real course questions before adding model complexity.")
roadmap = [
    ("0", "BASELINE", "Corpus + query set", "Record what fails today. Include text, tables, labeled diagrams, and scans if they occur.", BLUE),
    ("1", "STRUCTURE", "Parser + page map", "Adopt a structured parser. Preserve headings, blocks, tables, pages, and figure assets.", TEAL),
    ("2", "DESCRIBE", "Figure enrichment", "Add optional OCR and captions; track provenance and cache unchanged outputs.", GOLD),
    ("3", "RETRIEVE", "Evidence bundle", "Return text + figure hits with page citations and a viewable source reference.", CORAL),
]
for i, (n, stage, title_v, body, accent) in enumerate(roadmap):
    y = 2.36 + i*1.06
    ellipse(s, .92, y+.12, .42, .42, accent)
    text(s, 1.03, y+.23, .2, .1, n, 9, NAVY, True, align=PP_ALIGN.CENTER)
    if i < 3:
        line(s, 1.13, y+.58, 1.13, y+1.0, "425661", 1.1)
    text(s, 1.65, y+.07, 1.52, .18, stage, 9, accent, True)
    text(s, 3.12, y+.02, 2.62, .28, title_v, 15, CREAM, True, FONT_HEAD)
    text(s, 5.65, y+.04, 6.42, .53, body, 11, "CAD5D5", False, spacing=1.04)
    if i < 3:
        line(s, 1.65, y+.82, 12.3, y+.82, "2A3B46", .65)
rect(s, 1.65, 6.68, 10.65, .25, "263A43")
text(s, 1.9, 6.74, 10.1, .12, "OPTIONAL EXPERIMENT  /  Add ColPali-style page retrieval only if captioned retrieval still misses visual questions.", 8, TEAL, True)
footer(s, 7)


# 08 — Evaluation
s = new_slide(CREAM)
header(s, "07  /  PROOF OF QUALITY", "Measure whether the right page shows up", "A retrieval system earns trust by surfacing inspectable evidence, including when it has no answer.", False)
rect(s, .82, 2.45, 7.48, 3.95, WHITE, line="D8D3C8")
text(s, 1.13, 2.75, 6.65, .24, "QUESTION TYPE", 9, MUTED_DARK, True)
text(s, 6.44, 2.75, 1.33, .24, "EXPECTED HIT", 9, MUTED_DARK, True, align=PP_ALIGN.RIGHT)
rows = [
    ("Paragraph fact", "Text block · page"),
    ("Table lookup", "Table · page"),
    ("Diagram label", "Figure record · page"),
    ("Arrow / sequence", "Description + image"),
    ("No answer in source", "No strong evidence"),
]
for i, (q, hit) in enumerate(rows):
    y = 3.23 + i*.57
    line(s, 1.13, y+.43, 7.94, y+.43, "E5E0D7", .7)
    text(s, 1.13, y, 3.78, .22, q, 11, INK, i == 3)
    text(s, 4.73, y, 3.04, .22, hit, 10, "397C6B" if i != 4 else "A35442", True, align=PP_ALIGN.RIGHT)
rect(s, 8.68, 2.45, 3.81, 1.78, NAVY)
label(s, 8.99, 2.76, "Track", TEAL, 1.2)
text(s, 8.99, 3.15, 3.03, .7, "Right source/page\nin top results", 17, CREAM, True, FONT_HEAD, spacing=1.02)
rect(s, 8.68, 4.49, 3.81, 1.91, "DDEAE5")
label(s, 8.99, 4.8, "Also inspect", "397C6B", 2)
text(s, 8.99, 5.17, 3.0, .84, "Evidence supports claim\nGenerated text is labeled\nNo-answer stays no-answer", 11, INK, True, FONT_HEAD, spacing=1.08)
text(s, .88, 6.65, 11.2, .18, "NO INVENTED METRICS  /  Establish the baseline first; compare against the same questions after each phase.", 9, MUTED_DARK, True)
footer(s, 8, False)


# 09 — Principles / close
s = new_slide()
header(s, "08  /  DESIGN PRINCIPLES", "Useful, local, and honest about evidence", "The goal is not a bigger model. It is a better path from question to the page that supports it.")
principles = [
    ("LOCAL BY DEFAULT", "Keep course files and derived assets on the user's machine unless remote processing is explicitly enabled.", TEAL),
    ("SOURCE BEFORE SUMMARY", "Always preserve the original page/figure. Captions help search; they do not replace evidence.", BLUE),
    ("PROVENANCE THROUGHOUT", "Carry source, page, content type, model, and parser versions from ingestion to answer.", GOLD),
    ("COMPLEXITY BY EVIDENCE", "Start with structure + captions. Add visual retrieval only when measured misses justify it.", CORAL),
]
for i, (title_v, body, accent) in enumerate(principles):
    x = .82 + (i % 2)*6.0
    y = 2.45 + (i // 2)*1.65
    rect(s, x, y, 5.55, 1.3, PANEL, line=PANEL2)
    rect(s, x+.25, y+.28, .08, .74, accent, radius=False)
    text(s, x+.53, y+.25, 4.6, .2, title_v, 9, accent, True)
    text(s, x+.53, y+.58, 4.62, .54, body, 11, CREAM, False, spacing=1.03)
rect(s, .82, 6.0, 11.55, .63, "20353B")
text(s, 1.12, 6.22, 10.9, .2, "NEXT  /  Run Phase 0 on a handful of real lecture PDFs — then build the page-aware parser slice.", 11, TEAL, True)
footer(s, 9)


prs.core_properties.title = "Study RAG — From Course Files to Visual Evidence"
prs.core_properties.subject = "Project overview and roadmap for diagram-aware local RAG"
prs.core_properties.author = "Study RAG project"
prs.core_properties.keywords = "RAG, study materials, diagrams, retrieval, roadmap"
prs.save(OUT)
print(f"Wrote {OUT} ({len(prs.slides)} slides)")
